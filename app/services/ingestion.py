"""
Индексация корпуса RAG (блок 5.5): файлы data/<категория>/… → документы LlamaIndex с
метаданными → IngestionPipeline (нарезка блока 5.4 → эмбеддинги блока 5.1 → QdrantVectorStore)
с SimpleDocumentStore и DocstoreStrategy.UPSERTS.

    service = IngestionService(get_settings())
    report = service.run()                 # весь корпус RAG__DATA_DIR: новые и изменённые файлы
    report = service.run(files=[path])     # один файл (POST /documents/upload)
    print(report.summary())                # «66 файлов: 0 new, 0 changed, 66 unchanged, 0 failed, 0 removed…»

Чтение — SimpleDirectoryReader с ридером на формат (READERS):
- .pdf  — PyMuPDFReader: документ на страницу, номер страницы — в метаданных page;
- .docx — DocxReader (docx2txt) из llama-index-readers-file;
- .html — HTMLTagReader: только <article> (статья без шапки, меню и подвала сайта), нет его —
          <main>, затем <body>;
- .md   — MarkdownReader: документ на раздел; YAML front matter — в метаданные, а не в текст;
- .txt  — как есть (учебный корпус data/rag-block-03).

Метаданные (file_metadata): file_name — имя файла; source — путь от корня корпуса
(billing/refunds_v2.docx), он же id документа; category — первая папка пути (data/billing/… →
billing; файл в корне — general); format; title, author и created_at — из самого документа
(свойства DOCX, метаданные PDF, <meta> HTML, front matter Markdown); last_modified — дата
изменения файла (stat); version — из имени файла (refunds_v2 → 2, regulation_v2.3 → 2.3); у
страницы PDF — page. Плюс метки модели эмбеддингов и нарезки: сменили модель или нарезку —
у документов другой hash, и следующий запуск переиндексирует их сам. В вектор метаданные не
идут (excluded_embed_metadata_keys — все ключи): эмбеддинг — только текст фрагмента, как в замерах
блока 5.4; заголовок и так стоит в начале текста статьи.

Инкрементальность: docstore (var/ingest/<коллекция>_docstore.json) хранит hash каждого документа
(текст + метаданные). Запуск сравнивает файлы с ним: не изменился — пропуск без эмбеддингов
(UPSERTS), изменился — старые фрагменты удаляются из Qdrant и пишутся новые, файла больше нет —
его фрагменты удаляются. Id фрагмента — uuid5 от id документа и номера: даже без docstore
(удалили var/) повторная запись перезаписывает те же точки, а не добавляет копии.

Файл, который не читается (битый PDF, DOCX не zip, нет текста — скан без текстового слоя),
переименовывается в <имя>.failed и пишется в лог ingest_file_failed; остальные индексируются.
Ошибки Qdrant и модели эмбеддингов — не вина файла: запуск прерывается, файлы не трогаются.
"""
from __future__ import annotations

import asyncio
import fnmatch
import os
import re
import threading
import time
import uuid
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from pathlib import Path, PurePosixPath
from typing import TYPE_CHECKING, Any

from llama_index.core import SimpleDirectoryReader
from llama_index.core.embeddings import BaseEmbedding
from llama_index.core.ingestion import DocstoreStrategy, IngestionPipeline
from llama_index.core.schema import Document
from llama_index.core.storage.docstore import SimpleDocumentStore
from llama_index.readers.file import HTMLTagReader, MarkdownReader, PyMuPDFReader
from llama_index.vector_stores.qdrant import QdrantVectorStore
from pydantic import PrivateAttr
from qdrant_client.http.exceptions import ResponseHandlingException, UnexpectedResponse
from qdrant_client.models import VectorParams

from app.observability.logging import get_logger
from app.services import chunking, embeddings
from app.services.rag_common import RAGError, qdrant_error

if TYPE_CHECKING:
    from llama_index.core.node_parser import NodeParser

    from app.core.config import Settings

log = get_logger()

FORMATS = {".pdf": "pdf", ".docx": "docx", ".html": "html", ".htm": "html", ".md": "md", ".txt": "txt"}
IGNORE_FILE = ".ragignore"
FAILED_SUFFIX = ".failed"
DEFAULT_CATEGORY = "general"
# Пространство имён uuid5 для id фрагментов: сменить — значит получить другие id и дубли в Qdrant.
NODE_NAMESPACE = uuid.UUID("3b8f6c1e-2d4a-4f7b-9c5e-8a1d0e6f4b27")
# Модели — то, что помогает ответить и сослаться (заголовок, автор, дата, редакция, страница);
# технические поля ей ни к чему. В вектор не идёт ничего, кроме текста.
LLM_EXCLUDED = ["source", "format", "last_modified", "embed_model", "chunking"]
_VERSION = re.compile(r"[_-]v(\d+(?:\.\d+)*)$", re.IGNORECASE)
_FRONT_MATTER = re.compile(r"\A---\s*\n(.*?)\n---\s*(?:\n|\Z)", re.DOTALL)
_PDF_DATE = re.compile(r"D:(\d{4})(\d{2})(\d{2})")


class IngestError(RAGError):
    """Корпус не прочитать целиком: нет папки, файл вне корпуса, нет пакета ридера."""


class EmptyDocumentError(ValueError):
    """В файле нет текста: скан без текстового слоя, пустой документ, HTML без статьи."""


# ---------------------------------------------------------------- эмбеддинги блока 5.1
class ProjectEmbedding(BaseEmbedding):
    """Модель эмбеддингов блока 5.1 для LlamaIndex: те же векторы, префиксы E5, нормирование и
    кеш, что у bare-metal-версии и у коллекции блока 5.2."""

    _service: Any = PrivateAttr()

    def __init__(self, service: embeddings.EmbeddingService, **kwargs: Any) -> None:
        super().__init__(model_name=service.label, embed_batch_size=64, **kwargs)
        self._service = service

    def _get_query_embedding(self, query: str) -> list[float]:
        return self._service.embed_query(query)

    def _get_text_embedding(self, text: str) -> list[float]:
        return self._service.embed_documents([text])[0]

    def _get_text_embeddings(self, texts: list[str]) -> list[list[float]]:
        return self._service.embed_documents(list(texts))

    async def _aget_query_embedding(self, query: str) -> list[float]:
        return await asyncio.to_thread(self._get_query_embedding, query)

    async def _aget_text_embedding(self, text: str) -> list[float]:
        return await asyncio.to_thread(self._get_text_embedding, text)


# ---------------------------------------------------------------- ридеры
class ArticleHTMLReader(HTMLTagReader):
    """HTML-страница сайта: текст статьи (<article>) без шапки, меню и подвала. Нет <article> —
    <main>, затем <body>: загруженная руками страница тоже прочитается."""

    TAGS = ("article", "main", "body")

    def load_data(self, file: Path, extra_info: dict | None = None) -> list[Document]:
        for tag in self.TAGS:
            docs = [doc for doc in HTMLTagReader(tag=tag).load_data(file, extra_info) if doc.text.strip()]
            if docs:
                return docs
        return []


class FrontMatterMarkdownReader(MarkdownReader):
    """MarkdownReader (документ на раздел «#…»), но YAML front matter («---\\ntitle: …\\n---») не
    попадает в текст первого раздела: его поля читает file_metadata."""

    def parse_tups(self, filepath: str, errors: str = "ignore", fs: Any = None) -> list[tuple[str | None, str]]:
        _, content = split_front_matter(Path(filepath).read_text(encoding="utf-8", errors=errors))
        if self._remove_hyperlinks:
            content = self.remove_hyperlinks(content)
        if self._remove_images:
            content = self.remove_images(content)
        return self.markdown_to_tups(content)


class PdfPageReader(PyMuPDFReader):
    """PyMuPDFReader с каноническим «import pymupdf» (PyMuPDF 1.24+): ридер из
    llama-index-readers-file импортирует старое имя fitz, и PyMuPDF 1.28 печатает на каждый
    запуск «The `fitz` API is deprecated». Документ — страница, её номер — metadata["source"],
    как у PyMuPDFReader; полный путь файла в метаданные не кладётся."""

    def load(self, file_path: Path | str, metadata: bool = True, extra_info: dict | None = None) -> list[Document]:
        with open_pdf(Path(file_path)) as pdf:
            return [Document(text=page.get_text(), metadata={**(extra_info or {}), "source": str(page.number + 1)})
                    for page in pdf]


def open_pdf(path: Path) -> Any:
    """PDF из байтов, а не по пути. pymupdf.open(path) на битом файле бросает FileDataError и держит
    файл открытым, пока живо исключение; Windows не даёт переименовать открытый файл в .failed."""
    import pymupdf

    return pymupdf.open(stream=path.read_bytes(), filetype="pdf")


def readers() -> dict[str, Any]:
    """Ридер на расширение: PDF — PyMuPDFReader (pymupdf), DOCX — DocxReader (docx2txt), HTML —
    HTMLTagReader (bs4), Markdown — MarkdownReader; все из llama-index-readers-file."""
    from llama_index.readers.file import DocxReader

    html = ArticleHTMLReader()
    return {".pdf": PdfPageReader(), ".docx": DocxReader(), ".html": html, ".htm": html,
            ".md": FrontMatterMarkdownReader()}


# ---------------------------------------------------------------- метаданные
def split_front_matter(text: str) -> tuple[dict[str, Any], str]:
    """YAML front matter Markdown-файла и текст после него."""
    import yaml

    text = text.replace("\r\n", "\n")
    match = _FRONT_MATTER.match(text)
    if not match:
        return {}, text
    try:
        data = yaml.safe_load(match.group(1)) or {}
    except yaml.YAMLError:          # блок front matter есть, но не YAML: в текст он всё равно не идёт
        data = {}
    return (data if isinstance(data, dict) else {}), text[match.end():]


def category_of(relative: PurePosixPath) -> str:
    """Категория — первая папка пути от корня корпуса: billing/refunds_v2.docx → billing."""
    return relative.parts[0] if len(relative.parts) > 1 else DEFAULT_CATEGORY


def version_of(name: str) -> str | None:
    """Редакция из имени файла: refunds_v2.docx → 2, support-regulation_v2.3.pdf → 2.3."""
    stem = name
    for suffix in (FAILED_SUFFIX, *FORMATS):
        stem = stem.removesuffix(suffix)
    match = _VERSION.search(stem)
    return match.group(1) if match else None


def _iso_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and value.strip():
        match = _PDF_DATE.match(value.strip())
        if match:
            return "-".join(match.groups())
        try:
            return date.fromisoformat(value.strip()[:10]).isoformat()
        except ValueError:
            return None
    return None


def _clean(value: Any) -> str | None:
    text = " ".join(str(value).split()) if value is not None else ""
    return text or None


def document_properties(path: Path) -> dict[str, str | None]:
    """Заголовок, автор и дата создания из самого документа (не из файловой системы)."""
    fmt = FORMATS.get(path.suffix.lower())
    if fmt == "docx":
        import docx

        core = docx.Document(str(path)).core_properties
        return {"title": _clean(core.title), "author": _clean(core.author), "created_at": _iso_date(core.created)}
    if fmt == "pdf":
        with open_pdf(path) as pdf:
            meta = pdf.metadata or {}
        return {"title": _clean(meta.get("title")), "author": _clean(meta.get("author")),
                "created_at": _iso_date(meta.get("creationDate"))}
    if fmt == "html":
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")

        def meta(*names: str) -> str | None:
            for name in names:
                tag = soup.find("meta", attrs={"name": name})
                if tag and tag.get("content"):
                    return str(tag["content"])
            return None

        heading = soup.find("h1") or soup.find("title")
        return {"title": _clean(heading.get_text()) if heading else None, "author": _clean(meta("author")),
                "created_at": _iso_date(meta("dcterms.created", "created", "date"))}
    if fmt == "md":
        front, _ = split_front_matter(path.read_text(encoding="utf-8", errors="ignore"))
        return {"title": _clean(front.get("title")), "author": _clean(front.get("author")),
                "created_at": _iso_date(front.get("created_at") or front.get("date"))}
    return {}


def file_metadata(path: Path, root: Path) -> dict[str, Any]:
    """Метаданные файла — общие для всех его документов (страниц, разделов). Порядок ключей
    постоянный: от него зависит hash документа, а значит, «unchanged» при повторном запуске."""
    relative = PurePosixPath(path.relative_to(root).as_posix())
    props = document_properties(path)
    modified = datetime.fromtimestamp(path.stat().st_mtime, tz=UTC).astimezone().date().isoformat()
    meta = {"file_name": path.name, "source": relative.as_posix(), "category": category_of(relative),
            "format": FORMATS[path.suffix.lower()], "title": props.get("title") or path.stem,
            "author": props.get("author"), "created_at": props.get("created_at"), "last_modified": modified,
            "version": version_of(path.name)}
    return {key: value for key, value in meta.items() if value is not None}


def clean_text(text: str) -> str:
    """Переводы строк Windows, пробелы в концах строк, тройные пустые строки."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def read_file(path: Path, root: Path, *, extra: dict[str, Any] | None = None) -> list[Document]:
    """Документы одного файла с метаданными и id. Нет текста — EmptyDocumentError, файл не
    читается — исключение ридера (Exception с причиной в __cause__)."""
    fmt = FORMATS.get(path.suffix.lower())
    if fmt is None:
        raise IngestError(f"{path.name}: формат не поддерживается — нужны {', '.join(sorted(set(FORMATS)))}")
    base = {**file_metadata(path, root), **(extra or {})}
    raw = SimpleDirectoryReader.load_file(input_file=path, file_metadata=lambda _: {}, file_extractor=readers(),
                                          raise_on_error=True)
    documents = []
    for doc in raw:
        text = clean_text(doc.text)
        if not text:
            continue
        meta = dict(base)
        if fmt == "pdf":                 # PyMuPDFReader кладёт номер страницы в metadata["source"]
            meta["page"] = int(doc.metadata.get("source") or len(documents) + 1)
        documents.append(Document(text=text, metadata=meta))
    if not documents:
        raise EmptyDocumentError("нет текста (пустой документ или скан без текстового слоя)")
    source = base["source"]
    for number, doc in enumerate(documents, start=1):
        if fmt == "pdf":
            doc.id_ = f"{source}#p{doc.metadata['page']}"
        else:
            doc.id_ = source if len(documents) == 1 else f"{source}#{number}"
        doc.excluded_embed_metadata_keys = list(doc.metadata)
        doc.excluded_llm_metadata_keys = [key for key in LLM_EXCLUDED if key in doc.metadata]
    return documents


def file_of(doc_id: str) -> str:
    """Файл документа по его id: billing/x.pdf#p2 → billing/x.pdf."""
    return doc_id.split("#", 1)[0]


# ---------------------------------------------------------------- список файлов
def load_ignore(root: Path) -> list[str]:
    path = root / IGNORE_FILE
    if not path.is_file():
        return []
    lines = (line.strip() for line in path.read_text(encoding="utf-8").splitlines())
    return [line for line in lines if line and not line.startswith("#")]


def is_ignored(relative: str, patterns: Iterable[str]) -> bool:
    """Шаблоны .ragignore: «папка/» — всё внутри папки; со слешем — путь от корня корпуса;
    без слеша — имя файла на любой глубине."""
    name = relative.rsplit("/", 1)[-1]
    for pattern in patterns:
        if pattern.endswith("/"):
            if relative.startswith(pattern) or f"/{pattern}" in f"/{relative}":
                return True
        elif fnmatch.fnmatchcase(relative, pattern) or ("/" not in pattern and fnmatch.fnmatchcase(name, pattern)):
            return True
    return False


def list_files(root: Path) -> list[Path]:
    """Файлы корпуса: поддерживаемые форматы, без скрытых, .failed и того, что в .ragignore."""
    if not root.is_dir():
        raise IngestError(f"Нет папки корпуса {root}: проверьте RAG__DATA_DIR или путь в команде")
    patterns = load_ignore(root)
    found = []
    for path in root.rglob("*"):
        relative = path.relative_to(root).as_posix()
        if (path.is_file() and path.suffix.lower() in FORMATS
                and not any(part.startswith(".") for part in relative.split("/"))
                and not is_ignored(relative, patterns)):
            found.append(path)
    return sorted(found, key=lambda p: p.relative_to(root).as_posix())


def mark_failed(path: Path) -> Path | None:
    """x.pdf → x.pdf.failed: при следующем запуске файл не читается заново, а в папке видно, что с
    ним не так. Не вышло переименовать (файл открыт в Windows) — None, файл остаётся."""
    target = path.with_name(path.name + FAILED_SUFFIX)
    try:
        os.replace(path, target)
    except OSError as exc:
        log.warning("ingest_rename_failed", file=path.name, error=str(exc))
        return None
    return target


# ---------------------------------------------------------------- отчёт
STATES = ("new", "changed", "unchanged", "failed", "removed")


@dataclass
class FileResult:
    path: str                  # от корня корпуса: billing/refunds_v2.docx
    state: str                 # new | changed | unchanged | failed | removed | queued
    documents: int = 0         # документов LlamaIndex: страниц PDF, разделов Markdown
    error: str | None = None
    renamed_to: str | None = None


@dataclass
class IngestReport:
    collection: str
    root: str
    files: list[FileResult] = field(default_factory=list)
    chunks: int = 0            # фрагментов записано в Qdrant этим запуском
    seconds: float = 0.0

    def count(self, state: str) -> int:
        return sum(1 for item in self.files if item.state == state)

    @property
    def total(self) -> int:
        return sum(1 for item in self.files if item.state != "removed")

    def summary(self) -> str:
        counts = ", ".join(f"{self.count(state)} {state}" for state in STATES)
        return (f"{self.total} файлов: {counts}; фрагментов записано {self.chunks}; коллекция {self.collection}; "
                f"{self.seconds:.1f} с")

    def as_dict(self) -> dict[str, Any]:
        return {"collection": self.collection, "files": self.total, **{state: self.count(state) for state in STATES},
                "chunks": self.chunks, "seconds": round(self.seconds, 2)}


# ---------------------------------------------------------------- сервис
class IngestionService:
    """Индексатор корпуса: один на процесс (сервис) или на запуск (scripts/ingest.py).
    run() не запускается параллельно сам с собой (lock): старт сервиса и загрузка файла через
    API не пишут в docstore одновременно."""

    def __init__(self, cfg: Settings, *, client: Any = None, embed_model: BaseEmbedding | None = None,
                 data_dir: Path | None = None, collection: str | None = None) -> None:
        if client is None and not cfg.qdrant_url:
            raise RAGError("QDRANT_URL не задан: RAG хранит фрагменты в Qdrant — укажите адрес в .env")
        if cfg.embedding_dim is None:
            raise RAGError("EMBEDDING_DIM не задан: укажите размерность модели эмбеддингов в .env")
        self.cfg = cfg
        self.settings = cfg.rag
        self.root = Path(data_dir or cfg.rag.data_dir).resolve()
        self.collection = collection or cfg.rag.collection
        self.docstore_path = cfg.rag.docstore_dir / f"{self.collection}_docstore.json"
        if client is None:
            from app.services.vector_store import make_qdrant_client

            key = cfg.qdrant_api_key.get_secret_value() if cfg.qdrant_api_key else None
            client = make_qdrant_client(cfg.qdrant_url, key, asynchronous=False)
        self.client = client
        self.embed_model = embed_model or ProjectEmbedding(embeddings.default_service())
        self.chunking_label = chunking.label(cfg.rag.chunking, chunk_size=cfg.rag.chunk_size,
                                             chunk_overlap=cfg.rag.chunk_overlap,
                                             buffer_size=cfg.rag.semantic_buffer_size,
                                             breakpoint_percentile_threshold=cfg.rag.semantic_breakpoint)
        self.status: dict[str, FileResult] = {}      # последний результат по файлу — для GET /documents/status
        self.last_report: IngestReport | None = None
        self._lock = threading.Lock()
        self.docstore = self._load_docstore()

    # -------------------------------------------------------- состояние
    def _load_docstore(self) -> SimpleDocumentStore:
        if self.docstore_path.is_file():
            try:
                return SimpleDocumentStore.from_persist_path(str(self.docstore_path))
            except (ValueError, OSError) as exc:     # JSON обрезан: процесс убили во время записи
                log.warning("ingest_docstore_unreadable", path=self.docstore_path.name, error=str(exc),
                            note="начинаем с пустого: id фрагментов постоянные, дублей не будет")
        return SimpleDocumentStore()

    def _persist(self) -> None:
        self.docstore_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = self.docstore_path.with_suffix(".tmp")
        self.docstore.persist(str(temporary))
        os.replace(temporary, self.docstore_path)

    def known_documents(self) -> dict[str, dict[str, str]]:
        """Что в docstore: файл → {id документа: hash}."""
        known: dict[str, dict[str, str]] = {}
        for doc_hash, doc_id in self.docstore.get_all_document_hashes().items():
            known.setdefault(file_of(doc_id), {})[doc_id] = doc_hash
        return known

    def points(self) -> int:
        if not self.client.collection_exists(self.collection):
            return 0
        return self.client.count(self.collection, exact=True).count

    def _check_collection(self) -> None:
        """Коллекция есть — векторы той же длины, что у модели (EMBEDDING_DIM). Нет коллекции или
        она пуста, а docstore что-то помнит (удалили том Qdrant) — docstore сбрасывается, иначе
        UPSERTS счёл бы всё «unchanged» и коллекция осталась бы пустой."""
        if self.client.collection_exists(self.collection):
            vectors = self.client.get_collection(self.collection).config.params.vectors
            if isinstance(vectors, VectorParams) and vectors.size != self.cfg.embedding_dim:
                raise RAGError(f"Коллекция {self.collection}: векторы из {vectors.size} чисел, а EMBEDDING_DIM="
                               f"{self.cfg.embedding_dim} — её заполняла другая модель: python scripts/ingest.py --rebuild")
        if self.docstore.get_all_document_hashes() and not self.points():
            log.warning("ingest_docstore_reset", collection=self.collection,
                        note="коллекция пуста, а docstore помнит документы — индексируем заново")
            self.docstore = SimpleDocumentStore()

    # -------------------------------------------------------- конвейер
    def splitter(self) -> NodeParser:
        """Нарезка блока 5.4 (RAG__CHUNKING, по умолчанию recursive 512/64) с постоянными id
        фрагментов: uuid5 от id документа и номера фрагмента."""
        parser = chunking.build_splitter(
            self.settings.chunking, chunk_size=self.settings.chunk_size, chunk_overlap=self.settings.chunk_overlap,
            embed_model=self.embed_model, buffer_size=self.settings.semantic_buffer_size,
            breakpoint_percentile_threshold=self.settings.semantic_breakpoint)
        parser.id_func = stable_node_id
        return parser

    def pipeline(self) -> IngestionPipeline:
        store = QdrantVectorStore(client=self.client, collection_name=self.collection)
        return IngestionPipeline(transformations=[self.splitter(), self.embed_model], docstore=self.docstore,
                                 vector_store=store, docstore_strategy=DocstoreStrategy.UPSERTS)

    def label_documents(self, documents: Sequence[Document]) -> None:
        for doc in documents:
            doc.metadata["embed_model"] = self.embed_model.model_name
            doc.metadata["chunking"] = self.chunking_label
            doc.excluded_embed_metadata_keys = list(doc.metadata)
            doc.excluded_llm_metadata_keys = [key for key in LLM_EXCLUDED if key in doc.metadata]

    def resolve(self, path: Path) -> Path:
        """Путь файла внутри корпуса (для run(files=…)): вне корпуса id и категория не определить."""
        resolved = (path if path.is_absolute() else Path.cwd() / path).resolve()
        if not resolved.is_relative_to(self.root):
            raise IngestError(f"{path} — вне корпуса {self.root}: положите файл в data/<категория>/")
        return resolved

    # -------------------------------------------------------- запуск
    def run(self, files: Sequence[Path] | None = None, *, rebuild: bool = False) -> IngestReport:
        """files=None — весь корпус (и удаление фрагментов файлов, которых больше нет); список —
        только эти файлы. rebuild — удалить коллекцию и docstore и проиндексировать всё заново."""
        with self._lock:
            started = time.perf_counter()
            try:
                report = self._run(files, rebuild=rebuild)
            except (ResponseHandlingException, UnexpectedResponse) as exc:
                raise qdrant_error(exc, self.cfg.qdrant_url) from exc
            report.seconds = time.perf_counter() - started
            for item in report.files:
                self.status[item.path] = item
            self.last_report = report
            log.info("ingest_done", **report.as_dict())
            return report

    def _run(self, files: Sequence[Path] | None, *, rebuild: bool) -> IngestReport:
        if rebuild:
            if self.client.collection_exists(self.collection):
                self.client.delete_collection(self.collection)
            self.docstore = SimpleDocumentStore()
            self.docstore_path.unlink(missing_ok=True)
            log.info("ingest_rebuild", collection=self.collection)
        else:
            self._check_collection()
        paths = list_files(self.root) if files is None else [self.resolve(Path(p)) for p in files]
        report = IngestReport(self.collection, str(self.root))
        known = self.known_documents()
        documents: list[Document] = []
        stale: set[str] = set()
        seen: set[str] = set()
        for path in paths:
            relative = path.relative_to(self.root).as_posix()
            seen.add(relative)
            previous = known.get(relative, {})
            reason = None
            try:
                docs = read_file(path, self.root)
            except ImportError as exc:       # нет pymupdf или docx2txt — вина окружения, не файла
                raise IngestError(f"{relative}: {exc}. Установите зависимости: pip install -r requirements.txt") from exc
            except Exception as exc:  # noqa: BLE001 — любой сбой ридера — вина файла: .failed, остальные идут дальше
                reason = _reason(exc, path, relative)
            if reason is not None:
                # Переименование — уже вне except: traceback с кадрами ридера (и его открытыми
                # файлами) к этому моменту освобождён, иначе Windows не даёт переименовать файл.
                renamed = mark_failed(path) if path.exists() else None
                stale.update(previous)       # прежняя версия файла из индекса уходит вместе с ним
                log.warning("ingest_file_failed", file=relative, error=reason,
                            renamed_to=renamed.name if renamed else None)
                report.files.append(FileResult(relative, "failed", error=reason,
                                               renamed_to=renamed.relative_to(self.root).as_posix() if renamed else None))
                continue
            self.label_documents(docs)
            current = {doc.id_: doc.hash for doc in docs}
            if not previous:
                state = "new"
            elif previous == current:
                state = "unchanged"
            else:
                state = "changed"
                stale.update(set(previous) - set(current))     # страниц или разделов стало меньше
            if state != "unchanged":
                log.info("ingest_file", file=relative, state=state, documents=len(docs))
            report.files.append(FileResult(relative, state, documents=len(docs)))
            documents.extend(docs)
        if files is None:                    # полный запуск: файлы, которых больше нет в корпусе
            for relative in sorted(set(known) - seen):
                stale.update(known[relative])
                log.info("ingest_file", file=relative, state="removed")
                report.files.append(FileResult(relative, "removed", documents=len(known[relative])))
        if stale and self.client.collection_exists(self.collection):
            store = QdrantVectorStore(client=self.client, collection_name=self.collection)
            for doc_id in sorted(stale):
                store.delete(doc_id)
        for doc_id in stale:
            self.docstore.delete_document(doc_id, raise_error=False)
        try:
            nodes = self.pipeline().run(documents=documents) if documents else []
        except Exception:
            # Модель эмбеддингов или Qdrant упали посреди записи: UPSERTS уже удалил старые фрагменты
            # изменённых документов. Их hash из docstore убираем, чтобы следующий запуск записал их
            # заново, а не счёл «unchanged» при пустом месте в коллекции.
            pending = {item.path for item in report.files if item.state in ("new", "changed")}
            for doc in documents:
                if file_of(doc.id_) in pending:
                    self.docstore.delete_document(doc.id_, raise_error=False)
            self._persist()
            raise
        report.chunks = len(nodes)
        if nodes or stale or rebuild or not self.docstore_path.exists():
            self._persist()
        return report

    def queue(self, relative: str) -> None:
        """Файл принят к индексации (POST /documents/upload) — статус до окончания run()."""
        self.status[relative] = FileResult(relative, "queued")

    def close(self) -> None:
        self.client.close()


def stable_node_id(index: int, document: Any) -> str:
    """id_func нарезки: id фрагмента не меняется от запуска к запуску."""
    return str(uuid.uuid5(NODE_NAMESPACE, f"{document.id_}#{index}"))


def _reason(exc: BaseException, path: Path, relative: str) -> str:
    """Причина ошибки ридера: SimpleDirectoryReader заворачивает её в Exception("Error loading file").
    Полный путь (в Windows — с именем пользователя) в лог не идёт: только путь от корня корпуса."""
    cause = exc.__cause__ if str(exc) == "Error loading file" and exc.__cause__ else exc
    text = str(cause).strip()
    for full in (str(path), path.as_posix()):
        text = text.replace(full, relative)
    return f"{type(cause).__name__}: {text}"[:300] if text else type(cause).__name__
