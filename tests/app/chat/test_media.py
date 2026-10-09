"""
Медиа -> content-part (блок 4.3, app/chat/media.py) без сети и без модели.

- PDF -> {"type": "text", "text": "[документ PDF]:\\n..."}; DOCX — так же, с таблицами;
- PNG/JPEG -> {"type": "image_url", "image_url": {"url": "data:image/...;base64,..."}} —
  картинка уходит в основной chat.completions как есть, отдельного Vision-вызова нет;
- тип файла сверяется с первыми байтами, размер — с пределами MEDIA__*.

Образцы: samples/support_rules.pdf и .docx (кириллица, таблица — samples/_generate_documents.py),
samples/chart.png и samples/photo.jpg; PDF для крайних случаев собирается здесь же.
"""
from __future__ import annotations

import base64
from io import BytesIO
from pathlib import Path

import docx
import pytest
from fastapi import UploadFile
from starlette.datastructures import Headers

from app.chat.domain import MediaError
from app.chat.media import (
    DOCX_MIME,
    clean_text,
    extract_docx_text,
    extract_pdf_text,
    media_to_part,
    read_media,
)
from app.core.config import MediaSettings

ROOT = Path(__file__).resolve().parents[3]
SAMPLES = ROOT / "samples"


def upload(data: bytes, filename: str | None, content_type: str | None) -> UploadFile:
    headers = Headers({"content-type": content_type}) if content_type else Headers({})
    return UploadFile(file=BytesIO(data), filename=filename, headers=headers, size=len(data))


def make_pdf(pages: list[str]) -> bytes:
    """Минимальный PDF: по строке ASCII-текста на страницу, шрифт Helvetica."""
    objects: dict[int, bytes] = {1: b"<< /Type /Catalog /Pages 2 0 R >>",
                                 3: b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"}
    kids = []
    for number, text in enumerate(pages):
        page_id, content_id = 4 + 2 * number, 5 + 2 * number
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
        objects[content_id] = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)
        objects[page_id] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font "
                            f"<< /F1 3 0 R >> >> /Contents {content_id} 0 R >>").encode()
        kids.append(f"{page_id} 0 R")
    objects[2] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode()
    out, offsets = bytearray(b"%PDF-1.4\n"), {}
    for number in sorted(objects):
        offsets[number] = len(out)
        out += b"%d 0 obj\n%s\nendobj\n" % (number, objects[number])
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offsets[n]:010d} 00000 n \n".encode() for n in sorted(objects))
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def make_docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- PDF и DOCX
async def test_pdf_becomes_text_part():
    part = await media_to_part(upload(make_pdf(["Password reset link is valid for 30 minutes"]),
                                      "manual.pdf", "application/pdf"))
    assert part["type"] == "text"
    assert part["text"].startswith("[документ PDF]:\n")
    assert "valid for 30 minutes" in part["text"]
    assert set(part) == {"type", "text"}                         # готовый content-part OpenAI


async def test_russian_pdf_with_table():
    ref = await read_media(upload((SAMPLES / "support_rules.pdf").read_bytes(), "Регламент.pdf", "application/pdf"))
    assert (ref.kind, ref.mime, ref.filename) == ("document", "application/pdf", "Регламент.pdf")
    text = ref.part["text"]
    assert "Регламент технической поддержки «Личный кабинет»" in text
    assert "Не проходит" in text and "1 рабочий день" in text      # таблица приоритетов


async def test_docx_paragraphs_and_tables_in_order():
    data = make_docx("Привет", "Тарифы:", table=[["Тариф", "Цена"], ["Базовый", "0 ₽"], ["Про", "990 ₽"]])
    part = await media_to_part(upload(data, "тарифы.docx", DOCX_MIME))
    assert part == {"type": "text", "text": "[документ DOCX]:\nПривет\nТарифы:\nТариф | Цена\nБазовый | 0 ₽\nПро | 990 ₽"}


async def test_docx_sample_from_samples():
    text = extract_docx_text((SAMPLES / "support_rules.docx").read_bytes())
    assert "Высокий | Не проходит оплата картой | 1 час | 1 рабочий день" in text


async def test_document_with_generic_mime_is_detected_by_name_and_bytes():
    """Telegram иногда присылает документ как application/octet-stream."""
    ref = await read_media(upload(make_docx("Текст"), "notes.docx", "application/octet-stream"))
    assert (ref.kind, ref.mime) == ("document", DOCX_MIME)
    ref = await read_media(upload(make_pdf(["Hello"]), "a.pdf", None))
    assert (ref.kind, ref.mime) == ("document", "application/pdf")


async def test_long_document_is_cut_to_limit():
    words = " ".join(f"слово{i}" for i in range(10_000))            # ~90 000 символов
    part = await media_to_part(upload(make_docx(words), "long.docx", DOCX_MIME))
    body = part["text"].removeprefix("[документ DOCX]:\n")
    assert len(body) < 30_000 + 100 and body.endswith("символов]")
    assert "показано" in body.rsplit("\n", 1)[-1]


def test_pdf_page_limit_and_scan_heuristic():
    text = extract_pdf_text(make_pdf(["Page one text", "Page two text", "Page three text"]), max_pages=2)
    assert "Page two" in text and "Page three" not in text and "первые 2 из 3 страниц" in text
    with pytest.raises(MediaError) as caught:                     # 5 страниц и почти нет текста — скан
        extract_pdf_text(make_pdf(["a", "b", "c", "d", "e"]))
    assert caught.value.status == 422 and "скан" in caught.value.message
    with pytest.raises(MediaError) as caught:
        extract_pdf_text(make_pdf([""]))
    assert caught.value.code == "media_empty"


def test_broken_files_are_422_not_500():
    with pytest.raises(MediaError) as caught:
        extract_pdf_text(b"%PDF-1.4\n garbage")
    assert (caught.value.status, caught.value.code) == (422, "media_unreadable")
    with pytest.raises(MediaError) as caught:
        extract_docx_text(b"PK\x03\x04 not a docx")
    assert (caught.value.status, caught.value.code) == (422, "media_unreadable")


def test_clean_text_drops_invisible_characters():
    assert clean_text("Сброс\x00 па\u00adроля\u200b\r\n\n\n\nСсылка  \nконец") == "Сброс пароля\n\nСсылка\nконец"


# ---------------------------------------------------------------- картинки
@pytest.mark.parametrize("name, mime", [("chart.png", "image/png"), ("photo.jpg", "image/jpeg")])
async def test_image_becomes_data_url_part(name, mime):
    data = (SAMPLES / name).read_bytes()
    part = await media_to_part(upload(data, name, mime))
    assert part["type"] == "image_url" and set(part) == {"type", "image_url"}
    url = part["image_url"]["url"]
    prefix = f"data:{mime};base64,"
    assert url.startswith(prefix)
    assert base64.b64decode(url[len(prefix):], validate=True) == data          # пиксели как есть


async def test_image_mime_comes_from_bytes_not_from_client():
    """Telegram присылает фото как image/jpeg; PNG с таким заголовком уйдёт модели как PNG."""
    ref = await read_media(upload((SAMPLES / "chart.png").read_bytes(), "photo.jpg", "image/jpeg"))
    assert ref.mime == "image/png" and ref.part["image_url"]["url"].startswith("data:image/png;base64,")


# ---------------------------------------------------------------- отказы
@pytest.mark.parametrize("data, filename, content_type, status, code", [
    (b"%PDF-1.4 fake", "photo.jpg", "image/jpeg", 415, "media_unsupported"),          # не картинка внутри
    (b"not a pdf at all", "doc.pdf", "application/pdf", 422, "media_unreadable"),
    (b"\xd0\xcf\x11\xe0 old word", "old.doc", "application/msword", 415, "media_unsupported"),
    (b"plain text", "notes.txt", "text/plain", 415, "media_unsupported"),
    (b"", "empty.pdf", "application/pdf", 422, "media_unreadable"),
])
async def test_rejected_files(data, filename, content_type, status, code):
    with pytest.raises(MediaError) as caught:
        await read_media(upload(data, filename, content_type))
    assert (caught.value.status, caught.value.code) == (status, code)
    assert caught.value.message                                     # текст для пользователя


async def test_size_limits_are_413():
    limits = MediaSettings(max_image_bytes=2048)
    with pytest.raises(MediaError) as caught:
        await read_media(upload((SAMPLES / "chart.png").read_bytes(), "chart.png", "image/png"), settings=limits)
    assert (caught.value.status, caught.value.code) == (413, "media_too_large")
    assert "2 КБ" in caught.value.message


# ---------------------------------------------------------------- защита от «бомб» (ревью блока 4.3)
async def test_size_is_checked_before_reading_into_memory():
    limits = MediaSettings(max_image_bytes=2048, max_audio_bytes=2048, max_document_bytes=2048)

    class Unread(UploadFile):
        async def read(self, size: int = -1) -> bytes:            # pragma: no cover — не должен вызываться
            raise AssertionError("файл больше предела прочитан в память")

    declared = Unread(file=BytesIO(b""), filename="big.pdf", headers=Headers({"content-type": "application/pdf"}),
                      size=10_000)
    with pytest.raises(MediaError) as caught:
        await read_media(declared, settings=limits)
    assert (caught.value.status, caught.value.code) == (413, "media_too_large")
    unknown_size = UploadFile(file=BytesIO(b"%PDF-" + bytes(5000)), filename="big.pdf",
                              headers=Headers({"content-type": "application/pdf"}))
    with pytest.raises(MediaError) as caught:                      # размер неизвестен — читается не больше предела
        await read_media(unknown_size, settings=limits)
    assert caught.value.status == 413


def docx_bomb(xml_bytes: int) -> bytes:
    import zipfile

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", b" " * xml_bytes)
    return buffer.getvalue()


def test_docx_zip_bomb_is_refused_before_parsing():
    import time

    data = docx_bomb(11 * 1024 * 1024)
    assert len(data) < 100_000                                     # архив маленький, распакованный — 11 МБ
    started = time.perf_counter()
    with pytest.raises(MediaError) as caught:
        extract_docx_text(data)
    assert (caught.value.status, caught.value.code) == (413, "media_too_large")
    assert time.perf_counter() - started < 1


def test_zip_without_word_document_is_422():
    import zipfile

    buffer = BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "привет")
    with pytest.raises(MediaError) as caught:
        extract_docx_text(buffer.getvalue())
    assert (caught.value.status, caught.value.code) == (422, "media_unreadable")


def test_extraction_stops_when_enough_text():
    pdf = make_pdf([f"Page {n} " + "word " * 10 for n in range(10)])        # по ~60 символов на страницу
    text = extract_pdf_text(pdf, max_chars=60)
    assert "Page 1" in text and "Page 9" not in text
    docx_text = extract_docx_text(make_docx(*[f"Абзац {n} " + "слово " * 10 for n in range(50)]), max_chars=100)
    assert "Абзац 1 " in docx_text and "Абзац 49" not in docx_text


async def test_slow_parsing_times_out_and_keeps_its_slot(monkeypatch):
    import asyncio
    import time

    from app.chat import media

    def slow(*args) -> str:
        time.sleep(0.4)
        return "текст"

    monkeypatch.setattr(media, "extract_docx_text", slow)
    limits = MediaSettings(parse_timeout=0.05)
    with pytest.raises(MediaError) as caught:
        await read_media(upload(make_docx("x"), "a.docx", DOCX_MIME), settings=limits)
    assert caught.value.status == 422 and "слишком долго" in caught.value.message
    slots = media._slots()
    assert slots._value == media.PARSE_SLOTS - 1                   # поток ещё работает и держит место
    await asyncio.sleep(0.6)
    assert slots._value == media.PARSE_SLOTS                       # закончил — место свободно


def test_document_placeholder_has_no_filename():
    from app.chat.domain import MediaRef
    from app.chat.media import placeholder

    ref = MediaRef(kind="document", mime="application/pdf", size=1, filename="Договор Иванова.pdf",
                   part={"type": "text", "text": "x"})
    assert placeholder(ref) == "[документ PDF]"                    # копия текста попадает в превью лога
