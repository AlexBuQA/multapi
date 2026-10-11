"""
Корпус корпоративного RAG (блок 5.5): 66 документов «Личного кабинета» в data/<категория>/ в
четырёх форматах — HTML, DOCX, Markdown и PDF.

    python scripts/build_corpus.py               # перезаписать документы корпуса в data/
    python scripts/build_corpus.py --out tmp/    # собрать в другую папку (проверка, тесты)

Источник — база знаний проекта (блоки 5.1–5.2), а не выдуманные тексты:

- data/help_center.jsonl — 56 статей справочного центра: каждая статья — отдельный документ
  в папке своего раздела. Формат — как у настоящих источников компании: справка для
  пользователей — HTML-страницы сайта (account, start, notifications, mobile), документы
  финансового отдела и отдела внедрения — Word (billing, team, integrations), документация
  для разработчиков — Markdown (api), документы безопасности — PDF (data-management);
- статьи REG-1…REG-4 — регламент поддержки версии 2.3: один PDF, раздел на странице;
- data/faq.md — 22 частых вопроса: по файлу Markdown на раздел (faq-billing.md и т. д.);
- data/release_notes.md — 24 заметки о выпусках: PDF по полугодиям, заметка на странице.

У документов есть то, из чего ingestion (app/services/ingestion.py) берёт метаданные: автор и
дата создания — в свойствах DOCX, метаданных PDF, <meta> HTML и front matter Markdown;
редакция — в имени файла (refunds_v2.docx, support-regulation_v2.3.pdf). HTML-страница
сверстана как страница сайта: шапка, меню и подвал вокруг <article> — ридер берёт только статью.

Файлы детерминированы (фиксированные даты внутри DOCX и PDF): повторный запуск даёт те же
байты, git не видит изменений, а ingest — «0 changed». Файлы корпуса лежат в репозитории,
запускать скрипт нужно, только если поменялась база знаний. Нужны python-docx и pymupdf.
"""
from __future__ import annotations

import argparse
import html
import io
import sys
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, date, datetime, time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.services.documents import Chunk, read_jsonl, read_markdown

DATA = ROOT / "data"

# Папка (= категория в метаданных) по разделу статьи и формат документов раздела.
FOLDERS = {"account": "account", "start": "start", "notifications": "notifications", "mobile": "mobile",
           "billing": "billing", "team": "team", "integrations": "integrations", "api": "api",
           "data": "data-management", "support": "support"}
FORMATS = {"account": "html", "start": "html", "notifications": "html", "mobile": "html",
           "billing": "docx", "team": "docx", "integrations": "docx", "api": "md", "data-management": "pdf"}
TITLES = {"account": "Аккаунт и безопасность", "start": "Начало работы", "notifications": "Уведомления",
          "mobile": "Мобильное приложение", "billing": "Тарифы и оплата", "team": "Команда",
          "integrations": "Интеграции", "api": "API", "data-management": "Данные", "support": "Поддержка",
          "releases": "Что нового"}
AUTHORS = {"html": "Редакция справочного центра", "billing": "Финансовый отдел", "team": "Отдел внедрения",
           "integrations": "Отдел внедрения", "api": "Команда API",
           "data-management": "Служба информационной безопасности", "support": "Служба поддержки",
           "releases": "Продуктовая команда", "faq": "Служба поддержки"}

# Имена файлов статей. Суффикс _vN — редакция документа: прежние лежат в data/archive/
# (сроки возврата 2025 года, требования к паролю 2025 года, обмен с 1С 2025 года).
SLUGS = {
    "KB-001": "password-reset", "KB-002": "password-requirements_v2", "KB-003": "two-factor-auth",
    "KB-004": "change-email", "KB-005": "delete-account", "KB-006": "emails-not-arriving",
    "KB-007": "payment-documents", "KB-008": "api-keys", "KB-009": "api-limits",
    "KB-010": "mobile-login-problems", "KB-011": "registration", "KB-012": "login",
    "KB-013": "login-yandex-gosuslugi", "KB-014": "home-page", "KB-015": "language-and-timezone",
    "KB-016": "active-sessions", "KB-017": "login-history", "KB-018": "lost-authenticator",
    "KB-019": "change-password", "KB-020": "change-phone", "KB-021": "account-lock",
    "KB-022": "notification-types", "KB-023": "telegram-notifications", "KB-024": "push-notifications",
    "KB-025": "quiet-hours", "KB-026": "unsubscribe", "KB-027": "payment-methods", "KB-028": "auto-renewal",
    "KB-029": "change-tariff", "KB-030": "refunds_v2", "KB-031": "company-details",
    "KB-032": "card-payment-fails", "KB-033": "promo-codes", "KB-034": "api-key-permissions",
    "KB-035": "api-ip-allowlist", "KB-036": "webhooks", "KB-037": "api-versions", "KB-038": "api-error-401",
    "KB-039": "api-sandbox", "KB-040": "install-mobile-app", "KB-041": "biometric-login",
    "KB-042": "offline-mode", "KB-043": "invite-members", "KB-044": "member-roles",
    "KB-045": "transfer-ownership", "KB-046": "remove-member", "KB-047": "data-export",
    "KB-048": "data-import", "KB-049": "personal-data-protection", "KB-050": "restore-deleted-records",
    "KB-051": "1c-integration_v2", "KB-052": "deadline-calendar",
}
REGULATION = ("support-regulation_v2.3.pdf", "Регламент технической поддержки «Личный кабинет»",
              "Версия 2.3, действует с 1 сентября 2026 года.")
PRIORITIES = [
    ("Приоритет", "Пример", "Первый ответ", "Решение"),
    ("Критический", "Личный кабинет недоступен всем пользователям", "15 минут", "4 часа"),
    ("Высокий", "Не проходит оплата картой", "1 час", "1 рабочий день"),
    ("Средний", "Не приходит письмо для сброса пароля", "4 часа", "3 рабочих дня"),
    ("Низкий", "Вопрос о настройке уведомлений", "1 рабочий день", "5 рабочих дней"),
]
HALVES = {1: "первое полугодие", 2: "второе полугодие"}
FAQ_CREATED = "2026-06-15"
CSS = ("* {font-family: sans-serif; font-size: 11pt; line-height: 1.4;} h1 {font-size: 16pt;} "
       "h2 {font-size: 13pt;} td, th {border: 1px solid #888; padding: 3px;} .meta {color: #555; font-size: 9pt;}")
PAGE_RECT = (56, 56, 539, 786)        # A4 в пунктах, поля 2 см
ZIP_TIME = (2026, 1, 1, 0, 0, 0)      # время записей внутри DOCX: одинаковые байты при каждом запуске


@dataclass(frozen=True)
class Article:
    doc_id: str
    title: str
    body: str            # текст без заголовка
    created: str         # 2026-05-19
    category: str        # исходный раздел базы знаний (account, data, …)
    section: str = ""    # номер раздела справки (2.1)


def _date(chunk: Chunk) -> str:
    return chunk.created_at[:10]


def _ru_date(day: str) -> str:
    return date.fromisoformat(day).strftime("%d.%m.%Y")


def _version(slug: str) -> str | None:
    return slug.rsplit("_v", 1)[1] if "_v" in slug else None


def _meta_line(article: Article, slug: str) -> str:
    parts = [f"Раздел справки {article.section}", f"статья {article.doc_id}"]
    if version := _version(slug):
        parts.append(f"редакция {version} от {_ru_date(article.created)}")
    else:
        parts.append(f"обновлено {_ru_date(article.created)}")
    return " · ".join(parts)


def load_articles() -> tuple[list[Article], list[Article], list[Article]]:
    """Статьи справки, частые вопросы и заметки о выпусках из базы знаний data/."""
    import json

    sections = {}
    for line in (DATA / "help_center.jsonl").read_text(encoding="utf-8").splitlines():
        if line.strip():
            row = json.loads(line)
            sections[row["id"]] = row["section"]

    def articles(chunks: list[Chunk], *, split_title: bool) -> list[Article]:
        result = []
        for chunk in chunks:
            body = chunk.text.split("\n", 1)[1].strip() if split_title else chunk.text
            result.append(Article(chunk.doc_id, chunk.title, body, _date(chunk), chunk.category,
                                  sections.get(chunk.doc_id, "")))
        return result

    help_center = articles(read_jsonl(DATA / "help_center.jsonl", "help_center.jsonl"), split_title=False)
    faq = articles(read_markdown(DATA / "faq.md", "faq.md"), split_title=True)
    releases = articles(read_markdown(DATA / "release_notes.md", "release_notes.md"), split_title=True)
    return help_center, faq, releases


# ---------------------------------------------------------------- форматы
def write_html(path: Path, article: Article, slug: str, folder: str, author: str) -> None:
    esc = html.escape
    menu = "".join(f'<li><a href="/help/{name}/">{esc(title)}</a></li>'
                   for name, title in TITLES.items() if name != "releases")
    page = f"""<!DOCTYPE html>
<html lang="ru">
<head>
<meta charset="utf-8">
<title>{esc(article.title)} — Справка «Личный кабинет»</title>
<meta name="author" content="{esc(author)}">
<meta name="dcterms.created" content="{article.created}">
<meta name="description" content="Справочный центр «Личный кабинет», раздел «{esc(TITLES[folder])}»">
</head>
<body>
<header><a href="/help/">Справка «Личный кабинет»</a> › <a href="/help/{folder}/">{esc(TITLES[folder])}</a></header>
<nav><ul>{menu}</ul></nav>
<main>
<article id="{article.doc_id}">
<h1>{esc(article.title)}</h1>
<p>{esc(article.body)}</p>
<p class="meta">{esc(_meta_line(article, slug))}</p>
</article>
</main>
<footer>© 2026 «Личный кабинет». Не нашли ответ? Напишите в поддержку: support@example.com</footer>
</body>
</html>
"""
    path.write_text(page, encoding="utf-8", newline="\n")


def _fixed_zip(data: bytes) -> bytes:
    """DOCX — это zip: python-docx пишет в записи текущее время. Переписываем с постоянным."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            info = zipfile.ZipInfo(item.filename, ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            dst.writestr(info, src.read(item.filename))
    return out.getvalue()


def write_docx(path: Path, article: Article, slug: str, folder: str, author: str) -> None:
    import docx
    from docx.oxml.ns import qn

    document = docx.Document()
    for style in document.styles:        # шрифт с кириллицей во всех стилях абзацев
        if style.type == 1:
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts.set(qn(slot), "DejaVu Sans")
    document.add_heading(article.title, level=1)
    document.add_paragraph(article.body)
    document.add_paragraph(_meta_line(article, slug))
    created = datetime.combine(date.fromisoformat(article.created), time(), tzinfo=UTC)
    props = document.core_properties
    props.author = author
    props.last_modified_by = author
    props.title = article.title
    props.subject = f"Справка «Личный кабинет»: {TITLES[folder]}"
    props.category = folder
    props.keywords = article.doc_id
    props.created = created
    props.modified = created
    props.revision = int(_version(slug) or 1)
    props.comments = ""
    buffer = io.BytesIO()
    document.save(buffer)
    path.write_bytes(_fixed_zip(buffer.getvalue()))


def write_markdown(path: Path, title: str, author: str, created: str, extra: dict[str, str], body: str) -> None:
    import yaml

    # yaml.safe_dump, а не f-строки: «title: Частые вопросы: API» без кавычек — уже не YAML.
    front = yaml.safe_dump({"title": title, "author": author, "created_at": created, **extra},
                           allow_unicode=True, sort_keys=False, width=1000)
    path.write_text(f"---\n{front}---\n\n{body.strip()}\n", encoding="utf-8", newline="\n")


def write_pdf(path: Path, pages: list[str], *, title: str, author: str, created: str, subject: str) -> None:
    """PDF через PyMuPDF: страница — HTML-фрагмент. Шрифт с кириллицей встроен (Nimbus Sans из
    MuPDF), поэтому текст извлекается обратно без «кракозябр»."""
    import pymupdf

    document = pymupdf.open()
    for fragment in pages:
        page = document.new_page(width=595, height=842)
        _, scale = page.insert_htmlbox(pymupdf.Rect(*PAGE_RECT), fragment, css=CSS)
        if scale < 1:
            raise ValueError(f"{path.name}: текст не поместился на страницу")
    stamp = "D:" + created.replace("-", "") + "000000Z"
    document.set_metadata({"title": title, "author": author, "subject": subject, "keywords": "",
                           "creator": "scripts/build_corpus.py", "producer": "PyMuPDF",
                           "creationDate": stamp, "modDate": stamp})
    document.subset_fonts()
    document.save(path, garbage=4, deflate=True, no_new_id=True)
    document.close()


# ---------------------------------------------------------------- сборка
def build(out: Path) -> list[Path]:
    help_center, faq, releases = load_articles()
    written: list[Path] = []

    def target(folder: str, name: str) -> Path:
        directory = out / folder
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / name
        written.append(path)
        return path

    esc = html.escape
    regulation = [a for a in help_center if a.category == "support"]
    for article in (a for a in help_center if a.category != "support"):
        folder = FOLDERS[article.category]
        fmt, slug = FORMATS[folder], SLUGS[article.doc_id]
        author = AUTHORS["html"] if fmt == "html" else AUTHORS[folder]
        path = target(folder, f"{slug}.{fmt}")
        if fmt == "html":
            write_html(path, article, slug, folder, author)
        elif fmt == "docx":
            write_docx(path, article, slug, folder, author)
        elif fmt == "md":
            write_markdown(path, article.title, author, article.created,
                           {"section": article.section, "id": article.doc_id},
                           f"# {article.title}\n\n{article.body}\n\n{_meta_line(article, slug)}")
        else:
            write_pdf(path, [(f"<h1>{esc(article.title)}</h1><p>{esc(article.body)}</p>"
                              f"<p class='meta'>{esc(_meta_line(article, slug))}</p>")],
                      title=article.title, author=author, created=article.created,
                      subject=f"Справка «Личный кабинет»: {TITLES[folder]}")

    # Регламент поддержки: страница на раздел, сроки — таблицей, как в samples/support_rules.
    name, title, version_line = REGULATION
    table = "".join("<tr>" + "".join(f"<{'th' if i == 0 else 'td'}>{esc(cell)}</{'th' if i == 0 else 'td'}>"
                                     for cell in row) + "</tr>" for i, row in enumerate(PRIORITIES))
    pages = []
    for number, article in enumerate(regulation, start=1):
        head = f"<h1>{esc(title)}</h1><p>{esc(version_line)}</p>" if number == 1 else ""
        body = article.body
        if article.doc_id == "REG-2":
            body = body.split(". ", 1)[0] + ":"
            pages.append(f"{head}<h2>{number}. {esc(article.title)}</h2><p>{esc(body)}</p><table>{table}</table>")
            continue
        pages.append(f"{head}<h2>{number}. {esc(article.title)}</h2><p>{esc(body)}</p>")
    write_pdf(target("support", name), pages, title=title, author=AUTHORS["support"],
              created=max(a.created for a in regulation), subject="Регламент поддержки, версия 2.3")

    # Частые вопросы: файл на раздел, вопрос — заголовок «##».
    by_category: dict[str, list[Article]] = defaultdict(list)
    for article in faq:
        by_category[FOLDERS[article.category]].append(article)
    for folder, items in sorted(by_category.items()):
        heading = f"Частые вопросы: {TITLES[folder].lower() if folder != 'api' else 'API'}"
        body = f"# {heading}\n\n" + "\n\n".join(f"## {a.title}\n\n{a.body}" for a in items)
        write_markdown(target(folder, f"faq-{folder}.md"), heading, AUTHORS["faq"], FAQ_CREATED,
                       {"questions": ", ".join(a.doc_id for a in items)}, body)

    # Заметки о выпусках: PDF на полугодие, заметка на странице.
    halves: dict[tuple[int, int], list[Article]] = defaultdict(list)
    for note in releases:
        year, month = int(note.created[:4]), int(note.created[5:7])
        halves[(year, 1 if month <= 6 else 2)].append(note)
    for (year, half), notes in sorted(halves.items()):
        title = f"Что нового в «Личном кабинете»: {HALVES[half]} {year} года"
        pages = []
        for index, note in enumerate(notes):
            head = f"<h1>{esc(title)}</h1>" if index == 0 else ""
            pages.append(f"{head}<h2>{esc(note.title)}</h2>"
                         f"<p class='meta'>Дата выпуска: {_ru_date(note.created)} · раздел «{TITLES[FOLDERS[note.category]]}»</p>"
                         f"<p>{esc(note.body)}</p>")
        write_pdf(target("releases", f"release-notes_{year}-h{half}.pdf"), pages, title=title,
                  author=AUTHORS["releases"], created=max(n.created for n in notes), subject="Заметки о выпусках")
    return written


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Корпус блока 5.5 из базы знаний data/ (HTML, DOCX, MD, PDF)")
    parser.add_argument("--out", type=Path, default=DATA, help="куда писать (по умолчанию data/)")
    args = parser.parse_args(argv)
    written = build(args.out)
    by_format: dict[str, int] = defaultdict(int)
    for path in written:
        by_format[path.suffix.lstrip(".")] += 1
    total = sum(path.stat().st_size for path in written)
    print(f"Документов: {len(written)} в {len({p.parent for p in written})} папках, "
          + ", ".join(f"{fmt} — {count}" for fmt, count in sorted(by_format.items()))
          + f"; {total / 1024:.0f} КБ")
    return 0


if __name__ == "__main__":
    sys.exit(main())
