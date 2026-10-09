"""
Образцы документов для блока 4.3: регламент поддержки «Личного кабинета» в DOCX и PDF.

Их можно прислать боту и спросить по ним, а тесты (tests/app/chat/test_media.py) проверяют
на них разбор кириллицы и таблиц. Файлы коммитятся в репозиторий, поэтому повторный запуск
обычно не нужен.

    python samples/_generate_documents.py        # нужен python-docx; PDF — LibreOffice (soffice)
"""
from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Pt

SAMPLES = Path(__file__).resolve().parent
DOCX_PATH = SAMPLES / "support_rules.docx"
FONT = "DejaVu Sans"        # есть и в Linux, и в LibreOffice для Windows; Word подставит похожий

PRIORITIES = [
    ("Приоритет", "Пример", "Первый ответ", "Решение"),
    ("Критический", "Личный кабинет недоступен всем пользователям", "15 минут", "4 часа"),
    ("Высокий", "Не проходит оплата картой", "1 час", "1 рабочий день"),
    ("Средний", "Не приходит письмо для сброса пароля", "4 часа", "3 рабочих дня"),
    ("Низкий", "Вопрос о настройке уведомлений", "1 рабочий день", "5 рабочих дней"),
]


def build_docx() -> None:
    document = docx.Document()
    # Один шрифт с кириллицей для всех письменностей во всех стилях абзацев. Иначе LibreOffice
    # рисует кириллицу и пробелы разными шрифтами, и в тексте PDF слова склеиваются.
    for style in document.styles:
        if style.type == 1:                       # WD_STYLE_TYPE.PARAGRAPH
            fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
            for slot in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                fonts.set(qn(slot), FONT)
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading("Регламент технической поддержки «Личный кабинет»", level=1)
    document.add_paragraph("Версия 2.3, действует с 1 сентября 2026 года.")
    document.add_heading("1. Как обратиться в поддержку", level=2)
    document.add_paragraph("Заявку можно создать в Telegram-боте, в разделе «Помощь» личного кабинета или "
                           "письмом на support@example.com. В заявке укажите, что делали, что ожидали и что "
                           "произошло; приложите скриншот ошибки.")
    document.add_heading("2. Сроки ответа и решения", level=2)
    document.add_paragraph("Сроки считаются с момента регистрации заявки и зависят от приоритета:")
    table = document.add_table(rows=0, cols=4)
    table.style = "Table Grid"
    for row in PRIORITIES:
        cells = table.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            cell.text = text
    document.add_heading("3. Что поддержка не делает", level=2)
    document.add_paragraph("Поддержка не запрашивает и не присылает пароли и коды из СМС. Если вас просят "
                           "назвать код, это мошенники: завершите разговор и смените пароль.")
    document.add_heading("4. Эскалация", level=2)
    document.add_paragraph("Если срок первого ответа прошёл, напишите в бот «/status заявки» или позвоните "
                           "дежурному инженеру: 8 800 000-00-00 (круглосуточно).")
    document.save(DOCX_PATH)


def build_pdf() -> None:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        print("LibreOffice не найден: PDF не создан")
        return
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(SAMPLES), str(DOCX_PATH)],
                   check=True, capture_output=True, timeout=120)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    for path in (DOCX_PATH, DOCX_PATH.with_suffix(".pdf")):
        if path.exists():
            print(f"{path.name}: {path.stat().st_size} байт")
